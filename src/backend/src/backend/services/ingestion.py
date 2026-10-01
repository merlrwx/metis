import re
import unicodedata
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {DOCX_MIME_TYPE},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/x-markdown", "text/plain"},
    ".markdown": {"text/markdown", "text/x-markdown", "text/plain"},
}


class UnsupportedDocument(ValueError):
    pass


class DocumentTooLarge(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    metadata: dict


def clean_filename(filename: str) -> str:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    name = "".join(char for char in name if unicodedata.category(char) != "Cc")
    if not name or name in {".", ".."} or len(name) > 512:
        raise UnsupportedDocument("A valid file name is required")
    return name


def validate_upload(filename: str, declared_mime_type: str, data: bytes) -> str:
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentTooLarge("File exceeds the 20 MiB upload limit")
    name = clean_filename(filename)
    extension = PurePath(name).suffix.lower()
    accepted_types = MIME_TYPES.get(extension)
    mime_type = declared_mime_type.split(";", 1)[0].strip().lower()
    if accepted_types is None or mime_type not in accepted_types:
        raise UnsupportedDocument("Supported files are PDF, DOCX, TXT and Markdown")

    if extension == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF-"):
            raise UnsupportedDocument("File content does not match its PDF type")
        return "application/pdf"
    if extension == ".docx":
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                if (
                    "[Content_Types].xml" not in names
                    or "word/document.xml" not in names
                    or sum(info.file_size for info in archive.infolist())
                    > MAX_DOCX_UNCOMPRESSED_BYTES
                ):
                    raise UnsupportedDocument(
                        "File content does not match its DOCX type"
                    )
        except zipfile.BadZipFile as error:
            raise UnsupportedDocument(
                "File content does not match its DOCX type"
            ) from error
        return DOCX_MIME_TYPE

    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise UnsupportedDocument("Text files must use UTF-8 encoding") from error
    if "\x00" in text:
        raise UnsupportedDocument("Text files cannot contain binary data")
    return "text/markdown" if extension in {".md", ".markdown"} else "text/plain"


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\r\n", "\n")
    text = text.replace("\r", "\n").replace("\u00a0", " ")
    text = "".join(
        char
        for char in text
        if char in "\n\t" or not unicodedata.category(char).startswith("C")
    )
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(?m)^[ \t]+|[ \t]+$", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _with_offsets(
    parts: list[tuple[str | None, int | None, str]],
) -> tuple[str, list[dict]]:
    text_parts = []
    records = []
    offset = 0
    for title, level, content in parts:
        normalized = normalize_text(content)
        if not normalized:
            continue
        if text_parts:
            offset += 2
        start = offset
        text_parts.append(normalized)
        offset += len(normalized)
        records.append({"title": title, "level": level, "start": start, "end": offset})
    return "\n\n".join(text_parts), records


def _markdown_sections(text: str) -> list[tuple[str | None, int | None, str]]:
    sections: list[tuple[str | None, int | None, str]] = []
    current_title = None
    current_level = None
    lines = []
    for line in text.splitlines():
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            sections.append((current_title, current_level, "\n".join(lines)))
            current_title = heading.group(2).strip()
            current_level = len(heading.group(1))
            lines = [line]
        else:
            lines.append(line)
    sections.append((current_title, current_level, "\n".join(lines)))
    return sections


def _docx_sections(data: bytes) -> list[tuple[str | None, int | None, str]]:
    document = DocxDocument(BytesIO(data))
    sections: list[tuple[str | None, int | None, str]] = []
    title = None
    level = None
    content: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            value = block.text
            heading = re.fullmatch(r"Heading ([1-6])", block.style.name)
            if heading:
                sections.append((title, level, "\n".join(content)))
                title = value.strip() or None
                level = int(heading.group(1))
                content = [value]
            elif value.strip():
                content.append(value)
        elif isinstance(block, Table):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in block.rows
            ]
            if rows:
                content.append("\n".join(rows))
    sections.append((title, level, "\n".join(content)))
    return sections


def extract_document(filename: str, mime_type: str, data: bytes) -> ExtractedDocument:
    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf":
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedDocument("Encrypted PDFs are not supported")
        pages = []
        text_parts = []
        offset = 0
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = normalize_text(page.extract_text() or "")
            if text_parts:
                offset += 2
            start = offset
            text_parts.append(page_text)
            offset += len(page_text)
            pages.append({"page": page_number, "start": start, "end": offset})
        extracted_text = "\n\n".join(text_parts).strip()
        if not extracted_text:
            raise UnsupportedDocument("No extractable text was found")
        return ExtractedDocument(extracted_text, {"pages": pages})

    if extension == ".docx":
        sections = _docx_sections(data)
    else:
        text = data.decode("utf-8-sig")
        sections = (
            _markdown_sections(text)
            if extension in {".md", ".markdown"} or mime_type == "text/markdown"
            else [(None, None, text)]
        )
    extracted_text, records = _with_offsets(sections)
    if not extracted_text:
        raise UnsupportedDocument("No extractable text was found")
    return ExtractedDocument(extracted_text, {"sections": records})
