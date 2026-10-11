"""Deterministic synthetic source file builders for evaluation ingestion."""

import csv
import io
import re
import unicodedata

from backend.services import ingestion


def build_synthetic_document(
    text: str, format_name: str, section: str = "Synthetic evidence"
) -> tuple[str, str, bytes]:
    """Build a real supported file whose extracted text contains the supplied fact."""
    if format_name == "txt":
        return f"{_safe_name(section)}.txt", "text/plain", text.encode("utf-8")
    if format_name == "markdown":
        data = f"# {section}\n\n{text}".encode()
        return f"{_safe_name(section)}.md", "text/markdown", data
    if format_name == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(["Evidence section", section])
        writer.writerow(["Supporting information", text])
        return f"{_safe_name(section)}.csv", "text/csv", output.getvalue().encode()
    if format_name == "docx":
        from docx import Document

        document = Document()
        document.add_heading(section, level=1)
        for paragraph in text.split("\n\n"):
            document.add_paragraph(paragraph)
        output = io.BytesIO()
        document.save(output)
        return (
            f"{_safe_name(section)}.docx",
            ingestion.DOCX_MIME_TYPE,
            output.getvalue(),
        )
    if format_name == "pdf":
        return f"{_safe_name(section)}.pdf", "application/pdf", _pdf(text)
    raise ValueError(f"Unsupported evaluation format: {format_name}")


def _safe_name(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
    return slug[:64] or "synthetic-evidence"


def _pdf(text: str) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    normalized = unicodedata.normalize("NFKD", text).encode("ascii", "ignore")
    normalized = b" ".join(normalized.split())
    escaped = (
        normalized.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 10 Tf 50 740 Td (" + escaped + b") Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()
