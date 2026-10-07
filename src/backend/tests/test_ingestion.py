from io import BytesIO

import pytest
from backend.services import ingestion
from docx import Document


def test_pdf_extraction_preserves_page_offsets(sample_pdf):
    result = ingestion.extract_document("policy.pdf", "application/pdf", sample_pdf)

    assert result.text == "Medication incident policy"
    assert result.metadata == {"pages": [{"page": 1, "start": 0, "end": 26}]}


def test_markdown_extraction_normalizes_text_and_preserves_sections():
    result = ingestion.extract_document(
        "policy.md",
        "text/markdown",
        b"# Policy\r\n\r\n  Keep   records.\r\n## Review\r\nCheck weekly.",
    )

    assert result.text == "# Policy\n\nKeep records.\n\n## Review\nCheck weekly."
    assert result.metadata["sections"] == [
        {"title": "Policy", "level": 1, "start": 0, "end": 23},
        {"title": "Review", "level": 2, "start": 25, "end": 48},
    ]


def test_docx_extraction_keeps_heading_and_table_text():
    document = Document()
    document.add_heading("Incident policy", level=1)
    document.add_paragraph("Record the incident.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Nurse"
    output = BytesIO()
    document.save(output)

    result = ingestion.extract_document(
        "policy.docx", ingestion.DOCX_MIME_TYPE, output.getvalue()
    )

    assert result.text == "Incident policy\nRecord the incident.\nOwner | Nurse"
    assert result.metadata["sections"] == [
        {"title": "Incident policy", "level": 1, "start": 0, "end": 50}
    ]


def test_txt_extraction_normalizes_unicode_and_offsets():
    result = ingestion.extract_document(
        "notes.txt", "text/plain", b"Cafe\xcc\x81\r\n\r\nDone"
    )

    assert result.text == "Caf\u00e9\n\nDone"
    assert result.metadata["sections"] == [
        {"title": None, "level": None, "start": 0, "end": 10}
    ]


def test_upload_validation_rejects_mime_spoof_and_non_utf8_text():
    with pytest.raises(ingestion.UnsupportedDocument):
        ingestion.validate_upload("policy.pdf", "text/plain", b"%PDF-1.4")

    with pytest.raises(ingestion.UnsupportedDocument):
        ingestion.validate_upload("notes.txt", "text/plain", b"\xff")


def test_upload_validation_checks_pdf_and_docx_signatures(sample_pdf):
    assert (
        ingestion.validate_upload("policy.pdf", "application/pdf", sample_pdf)
        == "application/pdf"
    )
    with pytest.raises(ingestion.UnsupportedDocument, match="content"):
        ingestion.validate_upload("policy.pdf", "application/pdf", b"not a pdf")

    document = Document()
    document.add_paragraph("A Word document")
    output = BytesIO()
    document.save(output)
    assert (
        ingestion.validate_upload(
            "policy.docx", ingestion.DOCX_MIME_TYPE, output.getvalue()
        )
        == ingestion.DOCX_MIME_TYPE
    )


def test_extraction_rejects_documents_without_text(sample_pdf):
    result = ingestion.extract_document("policy.pdf", "application/pdf", sample_pdf)
    assert result.text
    with pytest.raises(ingestion.UnsupportedDocument, match="No extractable text"):
        ingestion.extract_document("notes.txt", "text/plain", b"  \n")


def test_upload_validation_rejects_oversized_file():
    with pytest.raises(ingestion.DocumentTooLarge):
        ingestion.validate_upload(
            "notes.txt", "text/plain", b"x" * (ingestion.MAX_UPLOAD_BYTES + 1)
        )


def test_csv_preserves_period_amount_and_headers_for_each_record():
    data = b'Period,Net AUD,Employee\nJanuary 2026,2000.10,"Smith, Jane"\nFebruary 2026,2100.25,"Smith, Jane"\n'
    assert ingestion.validate_upload("pay.csv", "text/csv", data) == "text/csv"
    result = ingestion.extract_document("pay.csv", "text/csv", data)
    from backend.services.chunking import split_document

    chunks = split_document(result.text, result.metadata)
    assert len(chunks) == 2
    assert chunks[0].section == "Row 1"
    assert "Period: January 2026\nNet AUD: 2000.10" in chunks[0].content
    assert "Employee: Smith, Jane" in chunks[1].content
    assert "Net AUD: 2100.25" in chunks[1].content


@pytest.mark.parametrize("data", [b"A,A\n1,2", b"A,B\n1", b"A,\n1,2"])
def test_csv_rejects_ambiguous_columns(data):
    with pytest.raises(ingestion.UnsupportedDocument):
        ingestion.extract_document("pay.csv", "text/csv", data)


def test_csv_enforces_row_and_cell_limits(monkeypatch):
    monkeypatch.setattr(ingestion, "MAX_CSV_ROWS", 1)
    with pytest.raises(ingestion.UnsupportedDocument, match="row limit"):
        ingestion.extract_document("pay.csv", "text/csv", b"A\n1\n2")
    with pytest.raises(ingestion.UnsupportedDocument, match="character limit"):
        ingestion.extract_document("pay.csv", "text/csv", b"A\n" + b"x" * 1001)


def test_scanned_pdf_has_actionable_recovery():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ingestion.UnsupportedDocument, match="Text could not be read"):
        ingestion.extract_document("scan.pdf", "application/pdf", output.getvalue())


@pytest.mark.parametrize(
    "title,headers,rows",
    [
        (
            "Fortnightly payslip",
            ["Pay period", "Net AUD", "Gross AUD"],
            [["1–14 January 2026", "2,000.10", "2,650.25"]],
        ),
        (
            "Invoice",
            ["Invoice number", "Due date", "Total AUD"],
            [["INV-2047", "31 January 2026", "1,025.50"]],
        ),
        (
            "Records policy",
            ["Record type", "Retention period"],
            [["Incident report", "7 years"], ["Payroll", "5 years"]],
        ),
    ],
)
def test_realistic_table_fixtures_preserve_labels_periods_and_numbers(
    title, headers, rows
):
    from backend.services.chunking import split_document

    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(headers))
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for row in rows:
        for cell, value in zip(table.add_row().cells, row, strict=True):
            cell.text = value
    output = BytesIO()
    document.save(output)
    result = ingestion.extract_document(
        "table.docx", ingestion.DOCX_MIME_TYPE, output.getvalue()
    )
    chunks = split_document(result.text, result.metadata)
    assert len(chunks) == 1
    assert chunks[0].section == title
    assert " | ".join(headers) in chunks[0].content
    for row in rows:
        assert ingestion.normalize_text(" | ".join(row)) in chunks[0].content


def test_pdf_with_only_page_number_is_not_ready(sample_pdf):
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(BytesIO(sample_pdf))
    page = reader.pages[0]
    stream = page["/Contents"].get_object()
    stream.set_data(stream.get_data().replace(b"Medication incident policy", b"Page 1"))
    writer = PdfWriter()
    writer.add_page(page)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ingestion.UnsupportedDocument, match="Text could not be read"):
        ingestion.extract_document(
            "nearly-empty.pdf", "application/pdf", output.getvalue()
        )


def test_pdf_table_keeps_pay_period_amount_labels_and_page_provenance(sample_pdf):
    from backend.services.chunking import split_document
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for period, net in [
        ("1-14 January 2026", "2,000.10"),
        ("15-28 January 2026", "2,100.25"),
    ]:
        page = PdfReader(BytesIO(sample_pdf)).pages[0]
        content = (
            f"BT /F1 12 Tf 72 720 Td (Pay period: {period}) Tj "
            f"0 -20 Td (Label | Amount AUD) Tj 0 -20 Td (Net pay | {net}) Tj "
            "0 -20 Td (Gross pay | 2,650.25) Tj ET"
        ).encode()
        page["/Contents"].get_object().set_data(content)
        writer.add_page(page)
    output = BytesIO()
    writer.write(output)
    result = ingestion.extract_document(
        "payslips.pdf", "application/pdf", output.getvalue()
    )
    chunks = split_document(result.text, result.metadata, max_tokens=254)
    assert len(chunks) == 2
    assert [chunk.page for chunk in chunks] == [1, 2]
    assert "1-14 January 2026" in chunks[0].content
    assert "Net pay | 2,000.10" in chunks[0].content
    assert "15-28 January 2026" in chunks[1].content
    assert "Net pay | 2,100.25" in chunks[1].content
    assert all("Label | Amount AUD" in chunk.content for chunk in chunks)
